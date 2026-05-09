from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.run import Run


def _set_font(
    run: Run,
    name: str = "Calibri",
    size: int = 11,
    bold: bool = False,
    color: tuple[int, int, int] | None = None,
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _get_field(section: dict | object, field: str) -> str:
    if isinstance(section, dict):
        return section[field]
    return getattr(section, field)


def build_word_document(sections: list) -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for i in range(1, 4):
        heading_style = doc.styles[f"Heading {i}"]
        heading_style.font.name = "Calibri"
        heading_style.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    _build_title_page(doc)
    _build_toc(doc, sections)
    _build_sections(doc, sections)
    _build_footer(doc)

    return doc


def _build_title_page(doc: Document) -> None:
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PROPOSAL")
    _set_font(run, size=28, bold=True, color=(0x1A, 0x3C, 0x6E))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Prepared in Response to RFP Requirements")
    _set_font(run, size=14, color=(0x66, 0x66, 0x66))

    doc.add_page_break()


def _build_toc(doc: Document, sections: list) -> None:
    doc.add_heading("Table of Contents", level=1)
    for s in sections:
        toc_entry = doc.add_paragraph()
        run = toc_entry.add_run(_get_field(s, "title"))
        _set_font(run, size=11, color=(0x1A, 0x3C, 0x6E))

    doc.add_page_break()


def _build_sections(doc: Document, sections: list) -> None:
    for s in sections:
        title_text = _get_field(s, "title")
        content_text = _get_field(s, "content")

        doc.add_heading(title_text, level=1)

        for para_text in content_text.split("\n\n"):
            para_text = para_text.strip()
            if not para_text:
                continue
            if para_text.startswith("# "):
                doc.add_heading(para_text[2:], level=2)
            elif para_text.startswith("## "):
                doc.add_heading(para_text[3:], level=3)
            elif para_text.startswith("- ") or para_text.startswith("* "):
                for line in para_text.split("\n"):
                    line = line.lstrip("- *").strip()
                    if line:
                        doc.add_paragraph(line, style="List Bullet")
            else:
                doc.add_paragraph(para_text)

        doc.add_page_break()


def _build_footer(doc: Document) -> None:
    section = doc.sections[-1]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CONFIDENTIAL")
    _set_font(run, size=8, color=(0x99, 0x99, 0x99))
