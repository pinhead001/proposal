import pytest
from io import BytesIO
from unittest.mock import patch
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def _make_docx(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@patch("app.api.routes.upload.save_proposals")
@patch("app.api.routes.upload.clear_analysis")
def test_upload_docx(mock_clear, mock_save):
    content = _make_docx("Past proposal content here")
    response = client.post(
        "/upload-proposals",
        files=[("files", ("test.docx", content, "application/octet-stream"))],
    )
    assert response.status_code == 200
    data = response.json()
    assert data["file_count"] == 1
    assert "extracted_texts" not in data
    mock_save.assert_called_once()
    mock_clear.assert_called_once()


@patch("app.api.routes.upload.save_proposals")
@patch("app.api.routes.upload.clear_analysis")
def test_upload_multiple_files(mock_clear, mock_save):
    files = [
        ("files", ("a.docx", _make_docx("Proposal A"), "application/octet-stream")),
        ("files", ("b.docx", _make_docx("Proposal B"), "application/octet-stream")),
    ]
    response = client.post("/upload-proposals", files=files)
    assert response.status_code == 200
    assert response.json()["file_count"] == 2


def test_upload_unsupported_format():
    response = client.post(
        "/upload-proposals",
        files=[("files", ("test.txt", b"plain text", "text/plain"))],
    )
    assert response.status_code == 400


def test_upload_empty_docx():
    doc = DocxDocument()
    buf = BytesIO()
    doc.save(buf)
    response = client.post(
        "/upload-proposals",
        files=[("files", ("empty.docx", buf.getvalue(), "application/octet-stream"))],
    )
    assert response.status_code == 400


@patch("app.api.routes.upload.MAX_BYTES", 10)
def test_upload_file_too_large():
    content = _make_docx("x" * 100)
    response = client.post(
        "/upload-proposals",
        files=[("files", ("big.docx", content, "application/octet-stream"))],
    )
    assert response.status_code == 413
