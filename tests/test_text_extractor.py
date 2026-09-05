import pytest
from io import BytesIO
from docx import Document
from app.services.extraction.text_extractor import extract_text


def _make_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx():
    content = _make_docx("Hello world from docx")
    result = extract_text(content, "test.docx")
    assert "Hello world from docx" in result


def test_extract_unsupported_format():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(b"data", "test.txt")


def test_extract_empty_docx():
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    with pytest.raises(ValueError, match="empty"):
        extract_text(buf.getvalue(), "empty.docx")
